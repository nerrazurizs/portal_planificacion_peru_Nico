"""Export utilities for CSV, Excel, and PowerPoint."""

import io
from datetime import datetime
import streamlit as st
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE


# ── Dorel corporate colors ─────────────────────────────────────────────────
_C_PRIMARY = RGBColor(0x06, 0x5E, 0x8B)
_C_ACCENT = RGBColor(0x23, 0xCE, 0xD3)
_C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_C_LIGHT_GRAY = RGBColor(0x94, 0xA3, 0xB8)
_C_DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)
_C_DANGER = RGBColor(0xDC, 0x26, 0x26)
_C_WARNING = RGBColor(0xD9, 0x77, 0x06)
_C_SUCCESS = RGBColor(0x16, 0xA3, 0x4A)
_C_NEUTRAL = RGBColor(0x64, 0x74, 0x8B)
_C_CARD_BG = RGBColor(0xF1, 0xF5, 0xF9)      # slate-100, fondo de tarjeta KPI
_C_CARD_BORDER = RGBColor(0xE2, 0xE8, 0xF0)  # slate-200, borde sutil
_C_PANEL_BG = RGBColor(0xF8, 0xFA, 0xFC)     # slate-50, fondo panel diagnóstico

# Leading emoji used by modules.business_case._build_bc_diagnostics, mapped
# to a severity color. Rendered as a colored "●" instead of the emoji
# itself — emoji-font fallback rendering is inconsistent across
# PowerPoint/Windows setups, a plain colored bullet renders identically
# everywhere. Covers both the "️" (FE0F) variation-selector and bare forms,
# since the source text mixes both.
_DIAG_SEVERITY = {
    "🚨": _C_DANGER, "📉": _C_DANGER, "🚫": _C_DANGER, "🔴": _C_DANGER,
    "⚠️": _C_WARNING, "⚠": _C_WARNING, "🏷️": _C_WARNING, "🏷": _C_WARNING,
    "🔧": _C_WARNING,
    "✅": _C_SUCCESS, "📈": _C_SUCCESS,
    "💡": _C_PRIMARY, "📦": _C_PRIMARY, "🏬": _C_PRIMARY, "💲": _C_PRIMARY,
    "📊": _C_PRIMARY, "📍": _C_PRIMARY,
    "❓": _C_NEUTRAL, "📅": _C_NEUTRAL, "🔄": _C_NEUTRAL, "🚢": _C_NEUTRAL,
    "📋": _C_NEUTRAL,
}


def _split_diag_bullet(text):
    """Strip a known leading emoji from a diagnostic bullet and return
    (severity_color, rest_of_text)."""
    text = str(text)
    for _emo, _color in _DIAG_SEVERITY.items():
        if text.startswith(_emo + " "):
            return _color, text[len(_emo) + 1:]
    return _C_NEUTRAL, text


def _add_markdown_runs(paragraph, text, size, color, bold_color=None):
    """Split **bold** markdown (as used by _build_bc_diagnostics, written for
    st.markdown) into alternating normal/bold runs within one paragraph —
    same emphasis as the on-screen module instead of plain stripped text."""
    for _idx, _part in enumerate(text.split("**")):
        if not _part:
            continue
        _r = paragraph.add_run()
        _r.text = _part
        _r.font.size = size
        _is_bold = _idx % 2 == 1
        _r.font.bold = _is_bold
        _r.font.color.rgb = (bold_color or color) if _is_bold else color


def _flatten_diag_lines(diag_list):
    """Expand diagnostic entries (which may have embedded '\n' sub-lines,
    e.g. "Venta por canal:\n  - Retail: ...\n  - Etail: ...") into a flat
    list of (is_sub, severity_color, text) — one tuple per visual line."""
    flat = []
    for _line in diag_list or []:
        _color, _text = _split_diag_bullet(str(_line))
        _sub_lines = _text.split("\n")
        flat.append((False, _color, _sub_lines[0]))
        for _sub in _sub_lines[1:]:
            _sub_clean = _sub.strip().lstrip("- ").strip()
            if _sub_clean:
                flat.append((True, _color, _sub_clean))
    return flat


def _fit_diag_lines(flat_lines, box_height_pt, max_size=11, min_size=6, n_columns=1):
    """Pick the largest font size that fits flat_lines in box_height_pt and,
    only if even min_size doesn't fit (a genuinely long diagnostic), trims
    the list and reports how many lines were cut.

    n_columns>1 lets the caller lay the same lines out side by side (e.g. 2
    text columns) — n rows of *vertical* space then needs to hold roughly
    n_lines/n_columns rows instead of n_lines, which is what actually makes
    a 15-20 line diagnostic fit at a readable size instead of forcing it to
    the size floor.

    This is the actual guarantee that text never overflows the slide —
    python-pptx's MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE only sets a flag asking
    PowerPoint to shrink at *open* time; it's a hint some viewers/renderers
    don't honor, so it can't be the only thing standing between this text
    and an overflowing box.

    Returns (font_size_pt, lines_to_render, n_omitted).
    """
    def _weighted(lines):
        return sum(1.3 if (not _is_sub and len(_t) > 90) else 1.0 for _is_sub, _c, _t in lines)

    if not flat_lines:
        return max_size, flat_lines, 0

    for _size in range(max_size, min_size - 1, -1):
        _line_h = _size * 1.22 + 1.5
        _rows_needed = _weighted(flat_lines) / n_columns
        if _rows_needed * _line_h <= box_height_pt:
            return _size, flat_lines, 0

    # Doesn't fit even at min_size — keep as many as provably fit (reserving
    # one row's worth for the "+N más" note) instead of letting it overflow.
    _line_h = min_size * 1.22 + 1.5
    _budget_rows = max(1.0, box_height_pt / _line_h - 1)
    _budget_weighted = _budget_rows * n_columns
    _kept, _acc = [], 0.0
    for _item in flat_lines:
        _w = 1.3 if (not _item[0] and len(_item[2]) > 90) else 1.0
        if _acc + _w > _budget_weighted:
            break
        _kept.append(_item)
        _acc += _w
    return min_size, _kept, len(flat_lines) - len(_kept)


_MAX_CSV_ROWS = 500_000
_MAX_EXCEL_ROWS = 200_000


def _df_to_csv_bytes(df: pd.DataFrame) -> bytes:
    """Write DataFrame to CSV bytes using BytesIO (avoids giant intermediate string)."""
    buf = io.BytesIO()
    df.to_csv(buf, index=False, encoding="utf-8-sig")
    return buf.getvalue()


def download_buttons(df, prefix="data", excel_buffer=None):
    """Render CSV and Excel download buttons for a DataFrame.

    Caps large DataFrames to avoid Streamlit message-size failures:
    - CSV:   first 500k rows
    - Excel: first 200k rows
    """
    col1, col2 = st.columns(2)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    n_rows = len(df)

    with col1:
        if n_rows > _MAX_CSV_ROWS:
            df_csv = df.iloc[:_MAX_CSV_ROWS]
            csv_label = f"Descargar CSV ({_MAX_CSV_ROWS // 1_000}k / {n_rows:,} filas)"
        else:
            df_csv = df
            csv_label = "Descargar CSV"

        st.download_button(
            csv_label,
            _df_to_csv_bytes(df_csv),
            f"{prefix}_{timestamp}.csv",
            "text/csv",
            key=f"dl_csv_{prefix}",
        )
        if n_rows > _MAX_CSV_ROWS:
            st.caption(
                f"Dataset tiene {n_rows:,} filas — descarga limitada a "
                f"{_MAX_CSV_ROWS:,}. Aplica filtros para reducir."
            )

    with col2:
        excel_rows = min(n_rows, _MAX_EXCEL_ROWS)
        if n_rows > _MAX_EXCEL_ROWS:
            excel_note = f" ({_MAX_EXCEL_ROWS // 1_000}k / {n_rows:,} filas)"
        else:
            excel_note = ""

        if excel_buffer is None:
            excel_buffer = io.BytesIO()
            with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
                df.iloc[:_MAX_EXCEL_ROWS].to_excel(writer, index=False)

        # Leer todos los bytes con seek(0) para que st.download_button
        # reciba bytes crudos (más confiable que pasar el objeto BytesIO)
        excel_buffer.seek(0)
        excel_bytes = excel_buffer.read()

        st.download_button(
            f"Descargar Excel{excel_note}",
            excel_bytes,
            f"{prefix}_{timestamp}.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"dl_excel_{prefix}",
        )


def _fig_to_png_bytes(fig):
    """Convert a matplotlib or Plotly figure to PNG bytes (or pass through bytes
    that were already rendered upstream, e.g. business_case's PPT chart)."""
    if isinstance(fig, io.BytesIO):
        fig.seek(0)
        return fig
    img_stream = io.BytesIO()
    if hasattr(fig, "savefig"):
        # matplotlib Figure
        fig.savefig(img_stream, format="png", bbox_inches="tight", dpi=150)
    elif hasattr(fig, "to_image"):
        # Plotly Figure — requires kaleido
        img_stream.write(fig.to_image(format="png", width=1200, height=700, scale=2))
    else:
        raise TypeError(f"Unsupported figure type: {type(fig)}")
    img_stream.seek(0)
    return img_stream


def _add_title_bar(slide, prs, text, title_width=None):
    """Add a Dorel-styled title bar at the top of a slide.

    title_width narrows the title textbox (default: nearly full slide width)
    so a caller can place other content — e.g. badges — to its right, inside
    the same bar, without overlapping the title text.
    """
    # Blue background bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), prs.slide_width, Inches(0.7),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = _C_PRIMARY
    bar.line.fill.background()

    # Accent teal line at bottom of bar
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0.7), prs.slide_width, Inches(0.04),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = _C_ACCENT
    accent.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.12), title_width or Inches(12), Inches(0.5),
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = _C_WHITE


def generate_ppt(figures_dict, title="Reporte Dorel"):
    """Generate a PowerPoint presentation from a dict of {slide_title: figure}.

    Supports both matplotlib and Plotly figures (auto-detected).
    Returns a BytesIO buffer with the .pptx content.

    Generates widescreen 16:9 slides with Dorel corporate styling.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # Blank

    # ── Slide 1: Title slide ──
    slide = prs.slides.add_slide(blank_layout)

    # Large blue background (60% of slide height)
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), prs.slide_width, Inches(4.5),
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = _C_PRIMARY
    bg.line.fill.background()

    # Accent teal line
    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(4.5), prs.slide_width, Inches(0.06),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = _C_ACCENT
    accent.line.fill.background()

    # Title text (centered, large)
    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11.33), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = _C_WHITE
    p.alignment = PP_ALIGN.CENTER

    # Subtitle with date
    txSub = slide.shapes.add_textbox(Inches(1), Inches(5.2), Inches(11.33), Inches(0.8))
    tf_sub = txSub.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = f"Generado el {datetime.now().strftime('%d de %B %Y')}"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = _C_LIGHT_GRAY
    p_sub.alignment = PP_ALIGN.CENTER

    # Footer
    txFoot = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(11.33), Inches(0.5))
    tf_foot = txFoot.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "Portal de Planificacion Dorel Chile"
    p_foot.font.size = Pt(10)
    p_foot.font.color.rgb = _C_LIGHT_GRAY
    p_foot.alignment = PP_ALIGN.CENTER

    # ── Content slides ──
    for name, fig in figures_dict.items():
        slide = prs.slides.add_slide(blank_layout)

        # Title bar
        _add_title_bar(slide, prs, name)

        # Figure as image (full width for 16:9)
        img_stream = _fig_to_png_bytes(fig)
        slide.shapes.add_picture(
            img_stream,
            Inches(0.3), Inches(0.9),
            width=Inches(12.7),
        )

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


def _accion_color(accion):
    """Map an ACCION_RAW label (business_case classification) to a badge color."""
    a = str(accion).upper()
    if any(k in a for k in ("PAUSAR", "LIQUIDAR", "DESCONTINUAR")):
        return _C_DANGER
    if any(k in a for k in ("REDUCIR", "MONITOREAR", "OBSERVAR")):
        return _C_WARNING
    if any(k in a for k in ("OK", "COMPRAR")):
        return _C_SUCCESS
    return _C_NEUTRAL


# Bandas de MOI (meses de inventario) — alineadas al "Diccionario de Estados
# de Salud" del módulo business_case: <2m es riesgo de quiebre (no cobertura
# sana), 2-6m saludable, 6-12m sobre-inventario, >12m crítico. Se agrega el
# corte sub-1m para distinguir el quiebre inminente del riesgo. MOI==0 se
# deja neutro a propósito: es ambiguo (puede ser sin stock = quebrado, o
# stock sin rotación = sobre-stock muerto), y esa distinción la resuelven
# los bullets del diagnóstico, no un solo número.
def _moi_severity_color(moi):
    """Color de severidad para el MOI en su tarjeta KPI."""
    m = float(moi or 0)
    if m <= 0:
        return _C_NEUTRAL
    if m < 1:
        return _C_DANGER      # quiebre inminente
    if m < 2:
        return _C_WARNING     # riesgo de quiebre
    if m <= 6:
        return _C_SUCCESS     # cobertura sana
    if m <= 12:
        return _C_WARNING     # sobre-stock
    return _C_DANGER          # sobre-stock crítico


def _moi_status_word(moi):
    """Palabra de estado que acompaña al MOI en su tarjeta KPI (equivalente a
    la línea 'Compared to…' del dashboard de referencia)."""
    m = float(moi or 0)
    if m <= 0:
        return "Sin stock / rotación"
    if m < 1:
        return "Quiebre inminente"
    if m < 2:
        return "Riesgo de quiebre"
    if m <= 6:
        return "Cobertura sana"
    if m <= 12:
        return "Sobre-stock"
    return "Sobre-stock crítico"


def _add_kpi_card(slide, left, top, width, height, value, label, accent,
                  value_color=None, status=None):
    """Tarjeta KPI estilo dashboard BI: fondo claro, barra de acento a la
    izquierda, valor grande en negrita, etiqueta en mayúsculas debajo y una
    línea de estado opcional. Reemplaza los chips planos anteriores para dar
    el look de reporte ejecutivo que pide gerencia (S&OP)."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = _C_CARD_BG
    card.line.color.rgb = _C_CARD_BORDER
    card.line.width = Pt(0.75)
    card.shadow.inherit = False

    # Barra de acento vertical a la izquierda (color = severidad/categoría).
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.09), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    bar.shadow.inherit = False

    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.22)
    tf.margin_right = Inches(0.1)
    tf.margin_top = Pt(3)
    tf.margin_bottom = Pt(3)

    p_label = tf.paragraphs[0]
    p_label.text = str(label).upper()
    p_label.font.size = Pt(9)
    p_label.font.bold = True
    p_label.font.color.rgb = _C_NEUTRAL

    p_val = tf.add_paragraph()
    p_val.text = str(value)
    p_val.font.size = Pt(21)
    p_val.font.bold = True
    p_val.font.color.rgb = value_color or _C_DARK_TEXT
    p_val.space_before = Pt(1)

    if status:
        p_status = tf.add_paragraph()
        p_status.text = str(status)
        p_status.font.size = Pt(8.5)
        p_status.font.color.rgb = accent
        p_status.space_before = Pt(1)


def _add_chip(slide, left, top, width, text, color):
    """Add a small rounded-rectangle badge with centered white text."""
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.32))
    chip.fill.solid()
    chip.fill.fore_color.rgb = color
    chip.line.fill.background()
    tf = chip.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = _C_WHITE
    p.alignment = PP_ALIGN.CENTER


def generate_otb_ppt(pm_name, entries, line_stats):
    """Generate the Business Case PPT: a title/summary slide plus one slide per SKU.

    Args:
        pm_name: planner/PM name shown on the title slide.
        entries: list of dicts with keys line, accion, mix_type, sku, nombre,
            fig (Plotly figure or None), kpis (dict), diagnostics (list[str]).
        line_stats: {line_label: {"n_skus": int, "stock_mm": str, "moi": float}}

    Returns a BytesIO buffer with the .pptx content.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # ── Slide 1: Title + line summary ──
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(4.5),
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = _C_PRIMARY
    bg.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(4.5), prs.slide_width, Inches(0.06),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = _C_ACCENT
    accent.line.fill.background()

    txBox = slide.shapes.add_textbox(Inches(1), Inches(1.3), Inches(11.33), Inches(1.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Caso de Negocio — OTB"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = _C_WHITE
    p.alignment = PP_ALIGN.CENTER

    txSub = slide.shapes.add_textbox(Inches(1), Inches(2.6), Inches(11.33), Inches(0.6))
    p_sub = txSub.text_frame.paragraphs[0]
    p_sub.text = f"{pm_name} · Generado el {datetime.now().strftime('%d de %B %Y')}"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = _C_LIGHT_GRAY
    p_sub.alignment = PP_ALIGN.CENTER

    _y = Inches(3.5)
    for _line_label, _stats in (line_stats or {}).items():
        txLine = slide.shapes.add_textbox(Inches(1), _y, Inches(11.33), Inches(0.45))
        p_line = txLine.text_frame.paragraphs[0]
        p_line.text = (
            f"{_line_label} — {_stats.get('n_skus', 0)} SKUs · "
            f"Stock {_stats.get('stock_mm', '—')}"
        )
        p_line.font.size = Pt(16)
        p_line.font.color.rgb = _C_WHITE
        p_line.alignment = PP_ALIGN.CENTER
        _y += Inches(0.45)

    txFoot = slide.shapes.add_textbox(Inches(1), Inches(6.7), Inches(11.33), Inches(0.5))
    p_foot = txFoot.text_frame.paragraphs[0]
    p_foot.text = "Portal de Planificación Dorel"
    p_foot.font.size = Pt(10)
    p_foot.font.color.rgb = _C_LIGHT_GRAY
    p_foot.alignment = PP_ALIGN.CENTER

    # ── One slide per SKU ──
    for entry in entries:
        slide = prs.slides.add_slide(blank_layout)
        _sku = entry.get("sku", "")
        _nombre = entry.get("nombre", "")
        # Badges live inside the title bar now (to the right of the SKU
        # title), not on their own row below it — that frees a whole row of
        # vertical space for the content area, which is where it's needed.
        _add_title_bar(
            slide, prs, f"{_sku} — {_nombre}" if _nombre else str(_sku),
            title_width=Inches(7.2),
        )
        _accion = str(entry.get("accion", "OK"))
        _mix = str(entry.get("mix_type", ""))
        _add_chip(slide, Inches(7.9), Inches(0.19), Inches(1.9), _accion, _accion_color(_accion))
        if _mix:
            _add_chip(slide, Inches(9.95), Inches(0.19), Inches(1.9), _mix, _C_PRIMARY)

        # ── KPI cards row (full width) ──
        # Tarjetas estilo dashboard BI (número grande + etiqueta + estado),
        # no chips planos: es el "de un vistazo" que gerencia lee primero en
        # S&OP. La línea/área ya aparece en el primer bullet del diagnóstico,
        # así que no se repite aquí — esa fila entera queda para los KPIs.
        _moi_actual = entry.get("moi_actual")
        _stock_cd = entry.get("stock_cd_und")
        _stock_total = entry.get("stock_total_und")
        _card_top = Inches(0.85)
        _card_h = Inches(0.95)
        _card_w = Inches(4.11)
        _card_gap = Inches(0.2)
        _card_x = [Inches(0.3), Inches(4.61), Inches(8.92)]
        if _moi_actual is not None:
            _moi_col = _moi_severity_color(_moi_actual)
            _add_kpi_card(
                slide, _card_x[0], _card_top, _card_w, _card_h,
                f"{_moi_actual:.1f} meses", "MOI Actual", _moi_col,
                value_color=_moi_col, status=_moi_status_word(_moi_actual),
            )
        if _stock_cd is not None:
            _add_kpi_card(
                slide, _card_x[1], _card_top, _card_w, _card_h,
                f"{_stock_cd:,.0f} und", "Stock CD (Centro de Distribución)", _C_PRIMARY,
            )
        if _stock_total is not None:
            _add_kpi_card(
                slide, _card_x[2], _card_top, _card_w, _card_h,
                f"{_stock_total:,.0f} und", "Stock Total (CD + Tiendas)", _C_ACCENT,
                value_color=_C_PRIMARY,
            )

        # ── Charts (full width, side-by-side figure) ──
        _chart_top = Inches(1.95)
        _fig = entry.get("fig")
        if _fig is not None:
            try:
                img_stream = _fig_to_png_bytes(_fig)
                slide.shapes.add_picture(
                    img_stream, Inches(0.3), _chart_top, width=Inches(12.73),
                )
            except Exception:
                _add_no_fig_note(slide, top=_chart_top)
            finally:
                if hasattr(_fig, "savefig"):
                    # matplotlib figures must be closed explicitly or they
                    # leak across hundreds of SKUs in a single PPT export.
                    import matplotlib.pyplot as plt
                    plt.close(_fig)
        else:
            _add_no_fig_note(slide, top=_chart_top)

        # ── Diagnostic panel (full width, all points) ──
        # Panel con fondo claro + encabezado, y TODOS los puntos del
        # diagnóstico (mismo _build_bc_diagnostics que el módulo en pantalla,
        # sin resumir). Se reparten en 2 columnas y el tamaño de fuente se
        # auto-ajusta para que entren completos; con 2 columnas en el ancho
        # total prácticamente nunca hay que recortar.
        _diag = entry.get("diagnostics") or []
        if _diag:
            _panel_top = Inches(5.42)
            _panel_h = Inches(1.95)   # termina en 7.37"
            _panel = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), _panel_top, Inches(12.73), _panel_h,
            )
            _panel.fill.solid()
            _panel.fill.fore_color.rgb = _C_PANEL_BG
            _panel.line.color.rgb = _C_CARD_BORDER
            _panel.line.width = Pt(0.75)
            _panel.shadow.inherit = False

            # Encabezado del panel (barra de acento + título).
            _hdr = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(0.3), _panel_top, Inches(0.09), Inches(0.34),
            )
            _hdr.fill.solid()
            _hdr.fill.fore_color.rgb = _C_PRIMARY
            _hdr.line.fill.background()
            _hdr.shadow.inherit = False
            _txHdr = slide.shapes.add_textbox(
                Inches(0.5), _panel_top + Inches(0.02), Inches(6.0), Inches(0.3),
            )
            _p_hdr = _txHdr.text_frame.paragraphs[0]
            _p_hdr.text = "DIAGNÓSTICO"
            _p_hdr.font.size = Pt(10)
            _p_hdr.font.bold = True
            _p_hdr.font.color.rgb = _C_PRIMARY

            # Área de texto: debajo del encabezado, dentro del panel.
            _diag_text_top = _panel_top + Inches(0.36)
            _diag_text_h = _panel_h - Inches(0.44)
            _flat = _flatten_diag_lines(_diag)
            _diag_size, _flat_fit, _n_omitted = _fit_diag_lines(
                _flat, box_height_pt=(_diag_text_h / 914400) * 72,
                max_size=11, min_size=6, n_columns=2,
            )
            _sub_size = max(6, _diag_size - 1)
            if _n_omitted > 0:
                _flat_fit = _flat_fit + [
                    (True, _C_NEUTRAL, f"… +{_n_omitted} puntos más (ver detalle en el módulo)"),
                ]

            _split = len(_flat_fit) - len(_flat_fit) // 2
            _columns = [_flat_fit[:_split], _flat_fit[_split:]]
            _col_lefts = [Inches(0.55), Inches(6.85)]
            _col_width = Inches(6.1)
            for _col_lines, _col_left in zip(_columns, _col_lefts):
                if not _col_lines:
                    continue
                txDiag = slide.shapes.add_textbox(_col_left, _diag_text_top, _col_width, _diag_text_h)
                tf_diag = txDiag.text_frame
                tf_diag.word_wrap = True
                tf_diag.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                tf_diag.vertical_anchor = MSO_ANCHOR.TOP
                tf_diag.margin_top = Pt(1)
                tf_diag.margin_bottom = Pt(1)
                for _idx, (_is_sub, _color, _text) in enumerate(_col_lines):
                    _p = tf_diag.paragraphs[0] if _idx == 0 else tf_diag.add_paragraph()
                    if _is_sub:
                        _p.space_after = Pt(1)
                        _r_dot = _p.add_run()
                        _r_dot.text = "    ◦ "
                        _r_dot.font.size = Pt(_sub_size)
                        _r_dot.font.color.rgb = _C_NEUTRAL
                        _add_markdown_runs(_p, _text, Pt(_sub_size), _C_NEUTRAL, bold_color=_C_DARK_TEXT)
                    else:
                        _p.space_after = Pt(2)
                        _r_bullet = _p.add_run()
                        _r_bullet.text = "● "
                        _r_bullet.font.size = Pt(_diag_size)
                        _r_bullet.font.bold = True
                        _r_bullet.font.color.rgb = _color
                        _add_markdown_runs(_p, _text, Pt(_diag_size), _C_DARK_TEXT, bold_color=_color)

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


def _add_no_fig_note(slide, top=None):
    txNoFig = slide.shapes.add_textbox(Inches(0.3), top or Inches(1.05), Inches(12.73), Inches(0.5))
    p = txNoFig.text_frame.paragraphs[0]
    p.text = "Sin datos suficientes para graficar."
    p.font.size = Pt(11)
    p.font.color.rgb = _C_NEUTRAL
    p.font.italic = True

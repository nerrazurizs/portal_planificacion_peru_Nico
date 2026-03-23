"""Listado O&E — Overstock & Excess for inter-division sales.

Generates a styled Excel export listing critical/excess stock from the
main CDs with transfer pricing at 5% margin in PEN and USD.
"""

import io

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import COLORS, apply_pm_filter
from db.cache import cached_query as cq
from utils.filters import calcular_moi_ajustado, human_format, norm_cols

# ── Constants ──
TC_USD_PEN = 3.80  # Peru exchange rate
MARGEN_INTERDIV = 0.05  # 5% margin on sale price

# LINEA → English category for international divisions
CATEGORY_MAP = {
    "COCHES": "Strollers",
    "TRAVEL SYSTEM": "Travel Systems",
    "SILLA AUTO": "Car Seats",
    "ACCESORIOS": "Accessories",
    "NURSERY": "Nursery",
    "PISCINAS": "Pools",
    "MI PRIMER JUGUETE": "My First Toy",
    "ROPA EXTERIOR": "Outerwear",
    "ROPA INTERIOR": "Underwear",
    "MOCHILAS Y LONCHERAS": "Backpacks & Lunchboxes",
    "ALIMENTACION": "Feeding",
    "BOTES": "Boats",
    "SPA": "Spa",
    "CALCETINES": "Socks",
    "PANTUFLAS": "Slippers",
    "HIGIENE": "Hygiene",
    "ESTRUCTURAS": "Structures",
    "JUGUETES MUSICALES": "Musical Toys",
    "CALZADO": "Footwear",
    "TOP": "Tops",
    "OUTWEAR": "Outerwear",
    "NEW BORN": "Newborn",
    "EASY SET": "Easy Set Pools",
    "ROPA INTERIOR/ROPA INTERIOR": "Underwear",
    "LACTANCIA": "Breastfeeding",
    "ROPA DE CAMA": "Bedding",
    "BIBERONES": "Bottles",
    "EXTRACTOR DE LECHE": "Breast Pumps",
    "ACCESORIOS NURSERY": "Nursery Accessories",
    "AUTO": "Car Seats",
    "DORMITORIO": "Bedroom",
    "BIENESTAR": "Wellness",
}


def _build_excel(df: pd.DataFrame) -> bytes:
    """Build styled Excel matching Dorel O&E corporate format."""
    buf = io.BytesIO()

    export_cols = [
        "SKU", "DESCRIPTION", "CATEGORY", "LINEA", "SUBLINEA", "BRAND",
        "QTY", "COST PEN", "COST USD", "TOTAL PEN", "TOTAL USD", "Warehouse",
    ]
    df_out = df[export_cols].copy()

    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df_out.to_excel(writer, index=False, sheet_name="O&E Inventory", startrow=1)
        ws = writer.sheets["O&E Inventory"]

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(export_cols))
        title_cell = ws.cell(row=1, column=1, value="DOREL PERU TOP O&E INVENTORY")
        title_cell.font = Font(bold=True, size=12)

        header_fill = PatternFill(start_color="065E8B", end_color="065E8B", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        thin_border = Border(bottom=Side(style="thin", color="23CED3"))
        for col_idx in range(1, len(export_cols) + 1):
            cell = ws.cell(row=2, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        money_cols_idx = {
            export_cols.index("COST PEN") + 1,
            export_cols.index("TOTAL PEN") + 1,
        }
        usd_cols_idx = {
            export_cols.index("COST USD") + 1,
            export_cols.index("TOTAL USD") + 1,
        }
        for row in range(3, len(df_out) + 3):
            for col_idx in money_cols_idx:
                cell = ws.cell(row=row, column=col_idx)
                cell.number_format = 'S/#,##0'
            for col_idx in usd_cols_idx:
                cell = ws.cell(row=row, column=col_idx)
                cell.number_format = '$#,##0.0'

        for col_idx, col_name in enumerate(export_cols, 1):
            max_len = len(col_name) + 2
            for row in range(3, min(len(df_out) + 3, 50)):
                val = ws.cell(row=row, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)) + 2)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len, 40)

    return buf.getvalue()


def render_listado_oe(conn):
    """Render the O&E inter-division sales module."""
    st.markdown("## Listado O&E — Venta Inter-Division")
    st.caption(
        "Stock del CD principal (CDU. Simple Primera Principal) para oferta a otras "
        "divisiones Dorel. Solo productos importados. Precio con margen 5%, "
        f"tipo de cambio USD/PEN = {TC_USD_PEN}."
    )

    with st.spinner("Consultando stock CDs principales..."):
        df_raw = norm_cols(cq.stock_oe(conn))
        df_vpx = cq.ventas_mensual_precio(conn)
        df_scd = cq.stock_cd_diario_6m(conn)

    if df_raw.empty:
        st.warning("Sin datos de stock en CDs principales.")
        return

    df_raw = apply_pm_filter(df_raw)

    # Only imported products (exclude NACIONAL)
    if "PROCEDENCIA" in df_raw.columns:
        df_raw = df_raw[df_raw["PROCEDENCIA"].astype(str).str.strip().str.upper() != "NACIONAL"]
    if df_raw.empty:
        st.warning("Sin productos importados en el CD principal.")
        return

    for c in ["QTY", "STOCK_COSTO", "ULTIMO_COSTO"]:
        if c in df_raw.columns:
            df_raw[c] = pd.to_numeric(df_raw[c], errors="coerce").fillna(0)

    df_raw = calcular_moi_ajustado(df_raw, df_vpx, stock_cd_diario=df_scd)

    col_f1, col_f2, col_f3 = st.columns(3)
    with col_f1:
        solo_critico = st.checkbox(
            "Solo stock critico (MOI>=12 + antiguedad>=12m)",
            value=True, key="oe_solo_crit",
        )
    with col_f2:
        areas = sorted(df_raw["AREA"].dropna().unique()) if "AREA" in df_raw.columns else []
        sel_area = st.multiselect("Area", areas, key="oe_area")
    with col_f3:
        lineas = sorted(df_raw["LINEA"].dropna().unique()) if "LINEA" in df_raw.columns else []
        sel_linea = st.multiselect("Linea", lineas, key="oe_linea")

    df = df_raw.copy()
    if solo_critico:
        _moi_src = df["MOI_CLASIF"] if "MOI_CLASIF" in df.columns else (df["MOI"] if "MOI" in df.columns else pd.Series(0, index=df.index))
        _moi_col = pd.to_numeric(_moi_src, errors="coerce").fillna(0)
        _age_src = df["MESES_EN_CIA"] if "MESES_EN_CIA" in df.columns else pd.Series(0, index=df.index)
        _age_col = pd.to_numeric(_age_src, errors="coerce").fillna(0)
        df = df[(_moi_col >= 12) & (_age_col >= 12)]

    if sel_area:
        df = df[df["AREA"].isin(sel_area)]
    if sel_linea:
        df = df[df["LINEA"].isin(sel_linea)]

    if df.empty:
        st.info("Sin SKUs con los filtros seleccionados.")
        return

    df["COST PEN"] = np.where(
        df["ULTIMO_COSTO"] > 0,
        df["ULTIMO_COSTO"] / (1 - MARGEN_INTERDIV),
        df["STOCK_COSTO"] / df["QTY"].replace(0, np.nan) / (1 - MARGEN_INTERDIV),
    )
    df["COST PEN"] = df["COST PEN"].fillna(0).round(0)
    df["COST USD"] = (df["COST PEN"] / TC_USD_PEN).round(1)
    df["TOTAL PEN"] = (df["QTY"] * df["COST PEN"]).round(0)
    df["TOTAL USD"] = (df["QTY"] * df["COST USD"]).round(1)

    df["CATEGORY"] = df["LINEA"].map(CATEGORY_MAP).fillna(df["LINEA"])

    df = df.rename(columns={
        "SKU_PRODUCTO": "SKU",
        "DESCRIPTION": "DESCRIPTION",
        "MARCA": "BRAND",
        "WAREHOUSE": "Warehouse",
    })
    if "DESCRIPTION" not in df.columns:
        for _fb in ["NOM_PRODUCTO", "SKU_NOM_PRODUCTO"]:
            if _fb in df.columns:
                df["DESCRIPTION"] = df[_fb]
                break
        else:
            df["DESCRIPTION"] = df.get("SKU", "")
    if "Warehouse" not in df.columns:
        df["Warehouse"] = "CDU. Simple Primera Principal"

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("SKUs", f"{df['SKU'].nunique():,}")
    k2.metric("Unidades", f"{df['QTY'].sum():,.0f}")
    k3.metric("Total PEN", f"S/{human_format(df['TOTAL PEN'].sum())}")
    k4.metric("Total USD", f"${human_format(df['TOTAL USD'].sum())}")

    st.dataframe(
        df[["SKU", "DESCRIPTION", "CATEGORY", "LINEA", "SUBLINEA", "BRAND",
            "QTY", "COST PEN", "COST USD", "TOTAL PEN", "TOTAL USD", "Warehouse"]],
        use_container_width=True,
        height=500,
        column_config={
            "COST PEN": st.column_config.NumberColumn(format="S/%,.0f"),
            "COST USD": st.column_config.NumberColumn(format="$%,.1f"),
            "TOTAL PEN": st.column_config.NumberColumn(format="S/%,.0f"),
            "TOTAL USD": st.column_config.NumberColumn(format="$%,.1f"),
            "QTY": st.column_config.NumberColumn(format="%,d"),
        },
    )

    st.download_button(
        "Descargar Excel O&E",
        _build_excel(df),
        f"dorel_peru_oe_inventory_{pd.Timestamp.now().strftime('%Y%m%d')}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )

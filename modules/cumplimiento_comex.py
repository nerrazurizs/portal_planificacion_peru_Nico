"""Cumplimiento Comex — dias de demora de embarque (ETD) e ingreso a almacen.

Toma todas las compras de comex en transito (Transito/Recibido) y las que ya
ingresaron (Cerrado) y compara las fechas comprometidas en la PO contra la
fecha real (una vez ocurrida) o la estimacion vigente (mientras sigue en
proceso):

  - Embarque:        PO_FECHA_DELIVERY  vs  ETD vigente
                      (real una vez zarpado, estimado mientras no zarpa)
  - Ingreso Almacen:  PO_FECHA_INGRESO_ALMACEN_ESTIMADO  vs  ingreso a CD vigente
                      (real una vez recepcionado, estimado mientras no llega)

Fuente: db_supply.fct.ft_cubo_comex (fechas comprometidas/estimadas de la PO)
        + db_supply.fct.ft_compras (fechas reales de embarque/recepcion).
"""

from __future__ import annotations

import calendar
import io
from datetime import date

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from config import COLORS, dorel_layout, apply_pm_filter
from db.cache import cached_query as cq
from db.queries import _VCM, _INSTOCK_CD, _PROD
from utils.export import download_buttons
from utils.filters import human_format, norm_cols
from utils.ui_animations import lottie_spinner
from utils.ui_components import page_header

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

_DATE_COLS = [
    "PO_FECHA_DELIVERY", "PO_FECHA_EMBARQUE", "FECHA_CARGOREADY",
    "FECHA_ETD_REAL", "FECHA_ETD_ESTIMADO_VIGENTE",
    "FECHA_ETA_PUERTO", "PO_FECHA_INGRESO_ALMACEN_ESTIMADO",
    "FECHA_INGRESO_ALMACEN_ESTIMADO_VIGENTE",
    "FECHA_REAL_INGRESO_ALMACEN",
]

_NUM_COLS = ["CANTIDAD_PEDIDA", "CANTIDAD_ENTREGADA", "VALOR_TOTAL_PO", "VALORIZADO_MN"]

_COLS_EMBARQUE = [
    # Identificacion
    "PO", "SKU_PRODUCTO", "NOM_PRODUCTO",
    # Clasificacion
    "PROVEEDOR", "AREA", "LINEA", "SUBLINEA", "MIX_OFICIAL",
    # Estado importacion
    "ESTADO_IMPORTACION",
    # Fechas clave (cronologico: compromiso → cargo ready → embarque)
    "PO_FECHA_DELIVERY", "FECHA_CARGOREADY", "ETD_VIGENTE", "ETD_CONFIRMADO",
    # Cumplimiento proveedor: Cargo Ready vs PO Delivery
    "CUMPLIMIENTO_PROVEEDOR_ONTIME",
    # Cumplimiento embarque: Cargo Ready vs ETD
    "DIAS_CARGO_TO_ETD", "CUMPLIMIENTO_EMBARQUE_ONTIME",
    # Cumplimiento fecha PO: ETD vs PO Delivery
    "DIAS_DEMORA_EMBARQUE", "CUMPLIMIENTO_FECHA_ETD_PO",
    # Impacto negocio
    "VP_ACUMULADA_SOLES",
]

_COLS_ALMACEN = [
    "PO", "SKU_PRODUCTO", "NOM_PRODUCTO", "PROVEEDOR", "AREA", "LINEA", "SUBLINEA",
    "ESTADO_IMPORTACION", "PO_FECHA_INGRESO_ALMACEN_ESTIMADO", "ALMACEN_VIGENTE",
    "ALMACEN_CONFIRMADO", "DIAS_DEMORA_ALMACEN", "CUMPLIMIENTO_ALMACEN",
]

_COL_CONFIG = {
    "PO": st.column_config.TextColumn("PO", width="small"),
    "SKU_PRODUCTO": st.column_config.TextColumn("SKU", width="small"),
    "NOM_PRODUCTO": st.column_config.TextColumn("Producto", width="medium"),
    "PROVEEDOR": st.column_config.TextColumn("Proveedor", width="medium"),
    "AREA": st.column_config.TextColumn("Area", width="small"),
    "LINEA": st.column_config.TextColumn("Linea", width="small"),
    "SUBLINEA": st.column_config.TextColumn("Sublinea", width="small"),
    "MIX_OFICIAL": st.column_config.TextColumn("Mix", width="small"),
    "ESTADO_IMPORTACION": st.column_config.TextColumn("Estado", width="small"),
    "PO_FECHA_DELIVERY": st.column_config.DateColumn("PO Fecha Delivery", format="DD/MM/YYYY"),
    "FECHA_CARGOREADY": st.column_config.DateColumn("Cargo Ready", format="DD/MM/YYYY"),
    "ETD_VIGENTE": st.column_config.DateColumn("ETD Vigente", format="DD/MM/YYYY"),
    "ETD_CONFIRMADO": st.column_config.CheckboxColumn("ETD Real"),
    "CUMPLIMIENTO_PROVEEDOR_ONTIME": st.column_config.TextColumn("Cumplimiento Proveedor Ontime", width="medium"),
    "DIAS_CARGO_TO_ETD": st.column_config.NumberColumn("CR → ETD (dias)", format="%d"),
    "CUMPLIMIENTO_EMBARQUE_ONTIME": st.column_config.TextColumn("Cumplimiento Embarque Ontime", width="medium"),
    "DIAS_DEMORA_EMBARQUE": st.column_config.NumberColumn("Dias Demora ETD", format="%d"),
    "CUMPLIMIENTO_FECHA_ETD_PO": st.column_config.TextColumn("Cumplimiento Fecha ETD PO", width="medium"),
    "VP_ACUMULADA_SOLES": st.column_config.NumberColumn("Venta Perdida (S/)", format="S/ %,.0f"),
    "PO_FECHA_INGRESO_ALMACEN_ESTIMADO": st.column_config.DateColumn("PO Fecha Ing. Almacen Est.", format="DD/MM/YYYY"),
    "ALMACEN_VIGENTE": st.column_config.DateColumn("Ingreso Almacen Vigente", format="DD/MM/YYYY"),
    "ALMACEN_CONFIRMADO": st.column_config.CheckboxColumn("Ingreso Real"),
    "DIAS_DEMORA_ALMACEN": st.column_config.NumberColumn("Dias Demora", format="%d"),
    "CUMPLIMIENTO_ALMACEN": st.column_config.TextColumn("Cumplimiento", width="small"),
}


# ---------------------------------------------------------------------------
# VP acumulada por SKU (quiebre CD → hoy)
# ---------------------------------------------------------------------------

def _demand_window_for_today() -> tuple[str, str, int]:
    """Ventana de demanda VCM: 2 meses completos anteriores, excluyendo diciembre.

    Misma logica que el modulo Venta Perdida. Retorna (start, end, dias) como strings.
    """
    today = date.today()
    m, y = today.month, today.year
    months: list[tuple[int, int]] = []
    cur_m, cur_y = m - 1, y
    if cur_m == 0:
        cur_m, cur_y = 12, y - 1
    while len(months) < 2:
        if cur_m == 12:
            cur_m -= 1
            if cur_m == 0:
                cur_m, cur_y = 11, cur_y - 1
            continue
        months.append((cur_y, cur_m))
        cur_m -= 1
        if cur_m == 0:
            cur_m, cur_y = 12, cur_y - 1
    oldest, newest = months[1], months[0]
    d_start = date(oldest[0], oldest[1], 1)
    d_end = date(newest[0], newest[1], calendar.monthrange(newest[0], newest[1])[1])
    return str(d_start), str(d_end), (d_end - d_start).days + 1


def _build_vp_acum_sql(dias_ventana: int) -> str:
    """SQL que calcula VP acumulada por SKU desde el quiebre CD hasta hoy.

    Logica: demanda diaria promedio (VCM 2 meses) x precio x dias sin stock CD.
    Solo incluye SKUs con quiebre activo (stock CD = 0 hoy).
    Params (%s): demand_start, demand_end
    """
    return f"""
WITH
demand_sku AS (
    SELECT
        v.sku_producto,
        SUM(v.cantidad) / {dias_ventana}::FLOAT      AS daily_demand,
        CASE WHEN SUM(v.cantidad) > 0
             THEN SUM(v.neto) / SUM(v.cantidad)
             ELSE 0 END                              AS avg_price_vcm
    FROM {_VCM} v
    WHERE v.fecha >= %s AND v.fecha <= %s
      AND v.cantidad > 0
      AND v.sku_producto NOT LIKE '%%-PV'
    GROUP BY 1
),
prod_price AS (
    SELECT p.sku_producto, MAX(p.ultimo_costo) AS ultimo_costo
    FROM {_PROD} p
    WHERE p.sku_producto NOT LIKE '%%-PV'
    GROUP BY 1
),
cd_hist AS (
    SELECT fecha, sku_producto, stock_unidades
    FROM {_INSTOCK_CD}
    WHERE fecha >= DATEADD('day', -90, CURRENT_DATE())
      AND fecha < CURRENT_DATE()
),
cd_latest AS (
    SELECT sku_producto, stock_unidades AS stock_actual
    FROM cd_hist
    QUALIFY ROW_NUMBER() OVER (PARTITION BY sku_producto ORDER BY fecha DESC) = 1
),
last_positive AS (
    SELECT sku_producto, MAX(fecha) AS last_positive_date
    FROM cd_hist
    WHERE stock_unidades > 0
    GROUP BY 1
),
quiebre AS (
    SELECT
        cs.sku_producto,
        DATEADD('day', 1,
            COALESCE(lp.last_positive_date,
                     DATEADD('day', -90, CURRENT_DATE()))
        )                                             AS fecha_quiebre,
        GREATEST(0, DATEDIFF('day',
            DATEADD('day', 1,
                COALESCE(lp.last_positive_date,
                         DATEADD('day', -90, CURRENT_DATE()))
            ),
            CURRENT_DATE()
        ))                                            AS dias_quiebre
    FROM cd_latest cs
    LEFT JOIN last_positive lp ON cs.sku_producto = lp.sku_producto
    WHERE cs.stock_actual = 0
)
SELECT
    q.sku_producto,
    GREATEST(0,
        COALESCE(d.daily_demand, 0)
        * COALESCE(NULLIF(d.avg_price_vcm, 0), pp.ultimo_costo, 0)
        * q.dias_quiebre
    )                                                 AS vp_acumulada_soles
FROM quiebre q
LEFT JOIN demand_sku d    ON q.sku_producto = d.sku_producto
LEFT JOIN prod_price pp   ON q.sku_producto = pp.sku_producto
WHERE q.dias_quiebre > 0
"""


@st.cache_data(ttl=3600, show_spinner=False)
def _load_vp_acum(
    _conn_id: int,
    demand_start: str,
    demand_end: str,
    dias_ventana: int,
    _conn=None,
) -> pd.DataFrame:
    """VP acumulada por SKU (quiebre CD activo) — cache 1h."""
    sql = _build_vp_acum_sql(dias_ventana)
    try:
        df = pd.read_sql(sql, _conn, params=[demand_start, demand_end])
        return norm_cols(df)
    except Exception:
        return pd.DataFrame()


# ---------------------------------------------------------------------------
# Data loading + enrichment
# ---------------------------------------------------------------------------

def _load(conn) -> pd.DataFrame:
    with lottie_spinner("comex"):
        df = cq.cumplimiento_comex(conn)
    return df


def _cumplimiento_label(dias) -> str:
    if pd.isna(dias):
        return "Sin dato"
    if dias <= 0:
        return "A tiempo"
    if dias <= 7:
        return "Atraso leve"
    return "Atraso critico"


def _cumplimiento_color(label: str) -> str:
    return {
        "A tiempo": COLORS["status_on_track"],
        "Atraso leve": COLORS["status_at_risk"],
        "Atraso critico": COLORS["status_critical"],
        "Sin dato": COLORS["medium_gray"],
    }.get(label, COLORS["medium_gray"])


def _cargo_etd_label(dias) -> str:
    """Clasifica atraso de Cargo Ready respecto al ETD (CargoReady − ETD).

    Negativo o <= 15 = carga lista antes o poco despues del zarpe.
    """
    if pd.isna(dias):
        return "Sin dato"
    if dias <= 15:
        return "OK"
    if dias <= 30:
        return "Retraso en embarque"
    return "Retraso en embarque Critico"


def _proveedor_ontime_label(dias) -> str:
    """Clasifica cumplimiento del proveedor (CargoReady − PO_FECHA_DELIVERY).

    En fecha: cargo lista al menos 15 dias antes del delivery comprometido.
    """
    if pd.isna(dias):
        return "Sin dato"
    if dias <= -15:
        return "Proveedor en fecha"
    if dias > 30:
        return "Proveedor Fuera de Fecha Critica"
    return "Proveedor fuera de fecha"


def _enrich(df: pd.DataFrame) -> pd.DataFrame:
    """Compute ETD/ALMACEN vigente dates, confirmed flags, and delay days."""
    if df.empty:
        return df
    df = df.copy()

    for c in _DATE_COLS:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors="coerce")
    for c in _NUM_COLS:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)

    df["PROVEEDOR"] = df["NOM_PROVEEDOR_MAESTRA"].fillna(df["NOM_PROVEEDOR_COMEX"])
    df.loc[df["PROVEEDOR"].astype(str).str.strip() == "", "PROVEEDOR"] = df["NOM_PROVEEDOR_COMEX"]

    # --- Embarque: confirmado cuando comex ya registro el ETD real
    # (di_etd / di_fechaembarque). Mientras no se registre, se usa el
    # estimado vigente (revisado por comex), no la fecha real de zarpe.
    df["ETD_CONFIRMADO"] = df["FECHA_ETD_REAL"].notna()
    df["ETD_VIGENTE"] = df["FECHA_ETD_REAL"].where(df["ETD_CONFIRMADO"], df["FECHA_ETD_ESTIMADO_VIGENTE"])
    df["DIAS_DEMORA_EMBARQUE"] = (df["ETD_VIGENTE"] - df["PO_FECHA_DELIVERY"]).dt.days

    # --- Ingreso a almacen: confirmado solo si ya existe la fecha real de
    # recepcion en CD (independiente del estado del expediente comex).
    df["ALMACEN_CONFIRMADO"] = df["FECHA_REAL_INGRESO_ALMACEN"].notna()
    df["ALMACEN_VIGENTE"] = df["FECHA_REAL_INGRESO_ALMACEN"].where(
        df["ALMACEN_CONFIRMADO"], df["FECHA_INGRESO_ALMACEN_ESTIMADO_VIGENTE"]
    )
    df["DIAS_DEMORA_ALMACEN"] = (df["ALMACEN_VIGENTE"] - df["PO_FECHA_INGRESO_ALMACEN_ESTIMADO"]).dt.days

    df["CUMPLIMIENTO_FECHA_ETD_PO"] = df["DIAS_DEMORA_EMBARQUE"].apply(_cumplimiento_label)
    df["CUMPLIMIENTO_ALMACEN"] = df["DIAS_DEMORA_ALMACEN"].apply(_cumplimiento_label)

    # Dias de atraso de Cargo Ready respecto al ETD vigente (positivo = carga lista despues del zarpe)
    if "FECHA_CARGOREADY" in df.columns and "ETD_VIGENTE" in df.columns:
        df["DIAS_CARGO_TO_ETD"] = (df["FECHA_CARGOREADY"] - df["ETD_VIGENTE"]).dt.days
    else:
        df["DIAS_CARGO_TO_ETD"] = np.nan
    df["CUMPLIMIENTO_EMBARQUE_ONTIME"] = df["DIAS_CARGO_TO_ETD"].apply(_cargo_etd_label)

    # Cumplimiento proveedor: Cargo Ready vs PO Fecha Delivery (compromiso original)
    if "FECHA_CARGOREADY" in df.columns and "PO_FECHA_DELIVERY" in df.columns:
        df["DIAS_CR_VS_PO"] = (df["FECHA_CARGOREADY"] - df["PO_FECHA_DELIVERY"]).dt.days
    else:
        df["DIAS_CR_VS_PO"] = np.nan

    is_sales_order = (
        df["ESTADO_IMPORTACION"].str.upper().str.strip() == "SALES ORDER"
        if "ESTADO_IMPORTACION" in df.columns
        else pd.Series(False, index=df.index)
    )
    hoy = pd.Timestamp.now().normalize()
    target_date = df["PO_FECHA_DELIVERY"] - pd.Timedelta(days=15)
    target_str = target_date.dt.strftime("%d/%m/%Y").where(target_date.notna(), "sin fecha")

    # Proveedor aun tiene tiempo: FCR target esta en el futuro
    msg_en_tiempo = "Próximo FCR " + target_str
    # FCR target ya vencio: alerta
    msg_vencido = "⚠️ FCR debio ser " + target_str + ", hacer seguimiento al Proveedor"

    so_msg = np.where(target_date > hoy, msg_en_tiempo, msg_vencido)

    df["CUMPLIMIENTO_PROVEEDOR_ONTIME"] = np.where(
        is_sales_order,
        so_msg,
        df["DIAS_CR_VS_PO"].apply(_proveedor_ontime_label),
    )

    # Status simplificado para filtros/segmentacion
    df["STATUS_COMPRA"] = np.where(df["ALMACEN_CONFIRMADO"], "Ya Ingreso", "En Transito")

    return df


# ---------------------------------------------------------------------------
# UI helpers
# ---------------------------------------------------------------------------

def _kpi(label: str, value: str, color: str, subtitle: str = "", icon: str = "") -> None:
    sub_html = (
        f'<div style="font-size:0.8rem;color:#64748b;margin-top:2px;'
        f'font-weight:500;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">'
        f"{subtitle}</div>" if subtitle else ""
    )
    icon_html = f'<span style="font-size:1.1rem;margin-right:4px;">{icon}</span>' if icon else ""
    st.html(
        f"""<div style="background:#fff;padding:0.7rem 0.9rem;border-radius:10px;
        box-shadow:0 1px 4px rgba(0,0,0,0.07);border-left:4px solid {color};
        min-height:80px;display:flex;flex-direction:column;justify-content:center;">
        <div style="font-size:0.7rem;color:#94a3b8;text-transform:uppercase;
        letter-spacing:0.5px;font-weight:600;margin-bottom:4px;">
        {icon_html}{label}</div>
        <div style="font-size:1.15rem;font-weight:700;color:{color};
        white-space:nowrap;overflow:hidden;text-overflow:ellipsis;">{value}</div>
        {sub_html}
    </div>"""
    )


def _select_cols(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    present = [c for c in cols if c in df.columns]
    return df[present].copy() if present else df.copy()


def _apply_filters(
    df: pd.DataFrame,
    sel_status: list[str],
    sel_estado: list[str],
    sel_proveedor: list[str],
    sel_area: list[str],
    sel_linea: list[str],
    sel_sublinea: list[str],
    sel_marca: list[str],
) -> pd.DataFrame:
    out = df
    if sel_status:
        out = out[out["STATUS_COMPRA"].isin(sel_status)]
    if sel_estado:
        out = out[out["ESTADO_IMPORTACION"].isin(sel_estado)]
    if sel_proveedor:
        out = out[out["PROVEEDOR"].isin(sel_proveedor)]
    if sel_area:
        out = out[out["AREA"].isin(sel_area)]
    if sel_linea:
        out = out[out["LINEA"].isin(sel_linea)]
    if sel_sublinea:
        out = out[out["SUBLINEA"].isin(sel_sublinea)]
    if sel_marca:
        out = out[out["MARCA"].isin(sel_marca)]
    return out


def _chart_demora_por_proveedor(df: pd.DataFrame, dias_col: str, title: str) -> None:
    # Excluye outliers extremos (> 1 ano) del promedio: son errores de
    # digitacion de fecha en el origen (ej. ano mal tipeado), no demoras
    # reales, y un solo valor asi distorsiona el promedio de un proveedor
    # entero. Las tablas de detalle si los muestran sin filtrar.
    work = df.dropna(subset=[dias_col])
    work = work[work[dias_col].abs() <= 365]
    if work.empty or "PROVEEDOR" not in work.columns:
        st.info("Sin datos suficientes para graficar.")
        return

    grouped = (
        work.groupby("PROVEEDOR", as_index=False)[dias_col]
        .mean()
        .sort_values(dias_col, ascending=True)
        .tail(15)
    )
    colors = [
        COLORS["status_critical"] if v > 7 else COLORS["status_at_risk"] if v > 0 else COLORS["status_on_track"]
        for v in grouped[dias_col]
    ]

    fig = go.Figure(
        go.Bar(
            x=grouped[dias_col],
            y=grouped["PROVEEDOR"],
            orientation="h",
            marker_color=colors,
            text=[f"{v:+.0f}d" for v in grouped[dias_col]],
            textposition="auto",
            hovertemplate="<b>%{y}</b><br>Demora promedio: %{x:.1f} dias<extra></extra>",
        )
    )
    fig.update_layout(
        **dorel_layout(
            title=title,
            height=max(350, len(grouped) * 30 + 100),
            xaxis=dict(title="Dias de demora promedio"),
            yaxis=dict(title="", tickfont=dict(size=10)),
        )
    )
    st.plotly_chart(fig, use_container_width=True)


# ---------------------------------------------------------------------------
# Tab renderers
# ---------------------------------------------------------------------------

def _render_resumen(df: pd.DataFrame) -> None:
    n_total = len(df)
    n_transito = int((df["STATUS_COMPRA"] == "En Transito").sum())
    n_ingresado = int((df["STATUS_COMPRA"] == "Ya Ingreso").sum())

    emb_confirmado = df[df["ETD_CONFIRMADO"]]
    alm_confirmado = df[df["ALMACEN_CONFIRMADO"]]

    pct_cumple_emb = (
        (emb_confirmado["DIAS_DEMORA_EMBARQUE"] <= 0).mean() * 100 if not emb_confirmado.empty else 0
    )
    pct_cumple_alm = (
        (alm_confirmado["DIAS_DEMORA_ALMACEN"] <= 0).mean() * 100 if not alm_confirmado.empty else 0
    )
    prom_demora_emb = emb_confirmado["DIAS_DEMORA_EMBARQUE"].mean() if not emb_confirmado.empty else np.nan
    prom_demora_alm = alm_confirmado["DIAS_DEMORA_ALMACEN"].mean() if not alm_confirmado.empty else np.nan

    # Alerta: PO con fecha de delivery ya vencida y aun sin ETD confirmado
    hoy = pd.Timestamp.now().normalize()
    en_riesgo = df[
        (~df["ETD_CONFIRMADO"]) & df["PO_FECHA_DELIVERY"].notna() & (df["PO_FECHA_DELIVERY"] < hoy)
    ]

    r1c1, r1c2, r1c3 = st.columns(3)
    with r1c1:
        _kpi("Lineas Totales", f"{n_total:,}", COLORS["primary"], icon="\U0001F4E6")
    with r1c2:
        _kpi("En Transito", f"{n_transito:,}", COLORS["tertiary_blue"], icon="\U0001F6A2")
    with r1c3:
        _kpi("Ya Ingresadas", f"{n_ingresado:,}", COLORS["status_on_track"], icon="✅")

    r2c1, r2c2, r2c3 = st.columns(3)
    with r2c1:
        _kpi(
            "Cumplimiento Embarque", f"{pct_cumple_emb:.0f}%",
            COLORS["status_on_track"] if pct_cumple_emb >= 70 else COLORS["status_at_risk"],
            subtitle=f"Demora prom.: {prom_demora_emb:+.1f}d" if pd.notna(prom_demora_emb) else "Sin ETD confirmado",
            icon="\U0001F4C5",
        )
    with r2c2:
        _kpi(
            "Cumplimiento Almacen", f"{pct_cumple_alm:.0f}%",
            COLORS["status_on_track"] if pct_cumple_alm >= 70 else COLORS["status_at_risk"],
            subtitle=f"Demora prom.: {prom_demora_alm:+.1f}d" if pd.notna(prom_demora_alm) else "Sin ingreso confirmado",
            icon="\U0001F3E2",
        )
    with r2c3:
        _kpi(
            "POs Vencidas sin Embarcar", f"{len(en_riesgo):,}",
            COLORS["status_critical"] if len(en_riesgo) > 0 else COLORS["status_on_track"],
            subtitle="PO Fecha Delivery ya paso, sin ETD confirmado",
            icon="⚠️",
        )

    st.caption(
        "**Confirmado** = ya ocurrio (fecha real de embarque/ingreso a almacen). "
        "Mientras una compra sigue en transito se usa la estimacion vigente mas reciente "
        "(revisada por comex), no la fecha original de la PO."
    )

    st.markdown("---")
    col_left, col_right = st.columns(2)
    with col_left:
        _chart_demora_por_proveedor(df, "DIAS_DEMORA_EMBARQUE", "Demora Promedio de Embarque por Proveedor (Top 15)")
    with col_right:
        _chart_demora_por_proveedor(df, "DIAS_DEMORA_ALMACEN", "Demora Promedio de Ingreso a Almacen por Proveedor (Top 15)")


_MANUAL_COLS = [
    ("Columna", "Descripcion"),
    ("PO", "Numero de orden de compra."),
    ("SKU_PRODUCTO", "Codigo del producto."),
    ("NOM_PRODUCTO", "Nombre del producto segun maestra de producto."),
    ("PROVEEDOR", "Nombre del proveedor. Se toma de la maestra de producto; si no existe, se usa el nombre registrado en comex."),
    ("AREA", "Area del producto segun maestra."),
    ("LINEA", "Linea del producto segun maestra."),
    ("SUBLINEA", "Sublinea del producto segun maestra."),
    ("MIX_OFICIAL", "Mix oficial del producto segun maestra."),
    ("ESTADO_IMPORTACION", "Estado actual del expediente de importacion en comex. Valores: Transito / Recibido / Cerrado / Sales Order."),
    ("PO_FECHA_DELIVERY", "Fecha de entrega comprometida al origen del proveedor segun la PO."),
    ("FECHA_CARGOREADY", "Fecha en que la carga estuvo fisicamente lista para embarque (FCR). Se toma el maximo registrado por PO + SKU en ft_compras. Vacio = aun no registrada."),
    ("ETD_VIGENTE", "Fecha de zarpe vigente: real si el buque ya zarpo (di_etd / di_fechaembarque), estimada si sigue en transito (di_fechaembarqueestimada / po_fechaembarque)."),
    ("ETD_CONFIRMADO", "TRUE si el ETD vigente es la fecha real de zarpe confirmada. FALSE si es una estimacion mientras sigue en transito."),
    (
        "CUMPLIMIENTO_PROVEEDOR_ONTIME",
        "Cumplimiento del proveedor en tener la carga lista (FCR) respecto al compromiso de la PO.\n"
        "- Sales Order: FCR objetivo = PO_FECHA_DELIVERY - 15 dias. Si la fecha aun no llego muestra 'Proximo FCR [fecha]'; si ya vencio muestra alerta.\n"
        "- Otros estados: diferencia FECHA_CARGOREADY - PO_FECHA_DELIVERY.\n"
        "  <= -15 dias = Proveedor en fecha\n"
        "  -14 a 30 dias = Proveedor fuera de fecha\n"
        "  > 30 dias = Proveedor Fuera de Fecha Critica",
    ),
    ("DIAS_CARGO_TO_ETD", "Dias entre Cargo Ready y ETD vigente (FECHA_CARGOREADY - ETD_VIGENTE). Negativo = carga lista antes del zarpe (bueno). Positivo = carga lista despues del zarpe (problema)."),
    (
        "CUMPLIMIENTO_EMBARQUE_ONTIME",
        "Clasificacion basada en DIAS_CARGO_TO_ETD:\n"
        "  <= 15 dias = OK\n"
        "  16 a 30 dias = Retraso en embarque\n"
        "  > 30 dias = Retraso en embarque Critico",
    ),
    ("DIAS_DEMORA_EMBARQUE", "Dias entre ETD vigente y PO Fecha Delivery (ETD_VIGENTE - PO_FECHA_DELIVERY). Positivo = el zarpe se atraso respecto al compromiso original de la PO."),
    (
        "CUMPLIMIENTO_FECHA_ETD_PO",
        "Clasificacion basada en DIAS_DEMORA_EMBARQUE:\n"
        "  <= 0 dias = A tiempo\n"
        "  1 a 7 dias = Atraso leve\n"
        "  > 7 dias = Atraso critico",
    ),
    ("VP_ACUMULADA_SOLES", "Venta Perdida acumulada en soles (S/) desde el quiebre de stock en CD hasta hoy. Solo para SKUs con stock CD = 0 actualmente. Calculo: demanda diaria promedio (VCM 2 meses, excl. diciembre) x precio x dias sin stock en CD."),
]


def _manual_row_height(desc: str, col_chars: int = 80, pts_per_line: float = 11.0) -> float:
    """Estima alto de fila en puntos basado en saltos de linea y longitud del texto."""
    import math
    lines = sum(max(1, math.ceil(len(seg) / col_chars)) if seg else 1 for seg in desc.split("\n"))
    return max(14, lines * pts_per_line)


def _excel_embarque(df: pd.DataFrame) -> io.BytesIO:
    """Excel formateado: hoja Manual + hoja datos. Fuente 8, auto-ajuste, fechas dd/mm/yyyy."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    hdr_fill = PatternFill(start_color="065E8B", end_color="065E8B", fill_type="solid")
    hdr_font = Font(bold=True, size=8, color="FFFFFF")
    data_font = Font(size=8)
    center = Alignment(horizontal="center", vertical="center")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ── Hoja 1: Manual ──────────────────────────────────────────────────────
    ws_m = wb.active
    ws_m.title = "Manual"

    for ci, h in enumerate(_MANUAL_COLS[0], 1):
        cell = ws_m.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center

    _COL_WIDTHS_MANUAL = [28, 80]
    for ri, (col_name, desc) in enumerate(_MANUAL_COLS[1:], 2):
        ws_m.cell(row=ri, column=1, value=col_name).font = data_font
        cell_desc = ws_m.cell(row=ri, column=2, value=desc)
        cell_desc.font = data_font
        cell_desc.alignment = wrap_top
        ws_m.row_dimensions[ri].height = _manual_row_height(desc, col_chars=_COL_WIDTHS_MANUAL[1])

    for ci, width in enumerate(_COL_WIDTHS_MANUAL, 1):
        ws_m.column_dimensions[get_column_letter(ci)].width = width
    ws_m.row_dimensions[1].height = 16
    ws_m.freeze_panes = "A2"

    # ── Hoja 2: Datos ────────────────────────────────────────────────────────
    ws = wb.create_sheet("Cumplimiento Embarque")

    df_out = df.copy()
    for col in df_out.columns:
        if pd.api.types.is_datetime64_any_dtype(df_out[col]):
            df_out[col] = df_out[col].dt.strftime("%d/%m/%Y").where(df_out[col].notna(), "")

    headers = list(df_out.columns)

    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = center

    for ri, row_vals in enumerate(df_out.itertuples(index=False), 2):
        for ci, val in enumerate(row_vals, 1):
            cell = ws.cell(row=ri, column=ci, value=val if not pd.isna(val) else None)
            cell.font = data_font

    for ci, header in enumerate(headers, 1):
        col_letter = get_column_letter(ci)
        col_series = df_out.iloc[:, ci - 1].astype(str)
        raw_max = col_series.str.len().max()
        max_data = 0 if pd.isna(raw_max) else int(raw_max)
        width = max(len(str(header)), max_data) + 2
        ws.column_dimensions[col_letter].width = max(8, min(width, 50))

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _render_tabla_embarque(df: pd.DataFrame) -> None:
    st.info(
        "**Dias Demora** = ETD Vigente − PO Fecha Delivery. "
        "**CR → ETD** = Cargo Ready − ETD vigente (dias). Negativo = carga lista antes del zarpe (OK). 0-15 OK, 16-30 Retraso en embarque, >30 Critico. "
        "**ETD Real** marcado = fecha confirmada; sin marcar, es la estimacion vigente. "
        "**Venta Perdida** = VP acumulada en S/ desde el quiebre de stock CD hasta hoy."
    )
    if df.empty:
        st.success("No hay compras para mostrar con los filtros actuales.")
        return

    display = _select_cols(df, _COLS_EMBARQUE).sort_values("DIAS_DEMORA_EMBARQUE", ascending=False, na_position="last")
    st.dataframe(
        display, column_config=_COL_CONFIG, use_container_width=True, hide_index=True,
        height=min(len(display) * 35 + 60, 600),
    )
    n_atraso = int((display["DIAS_DEMORA_EMBARQUE"] > 0).sum())
    n_cargo_critico = int((display["CUMPLIMIENTO_EMBARQUE_ONTIME"] == "Retraso en embarque Critico").sum()) if "CUMPLIMIENTO_EMBARQUE_ONTIME" in display.columns else 0
    vp_total = (
        pd.to_numeric(display["VP_ACUMULADA_SOLES"], errors="coerce").sum()
        if "VP_ACUMULADA_SOLES" in display.columns else 0
    )
    vp_caption = f" | VP total: S/ {vp_total:,.0f}" if vp_total > 0 else ""
    cargo_caption = f" | {n_cargo_critico} Cargo Ready criticos" if n_cargo_critico > 0 else ""
    st.caption(f"{len(display)} lineas | {n_atraso} con atraso de embarque{cargo_caption}{vp_caption}")
    download_buttons(display, prefix="cumplimiento_embarque", excel_buffer=_excel_embarque(display))


def _render_tabla_almacen(df: pd.DataFrame) -> None:
    st.info(
        "**Dias Demora** = Ingreso Almacen Vigente − PO Fecha Ingreso Almacen Estimado. Valores "
        "positivos indican atraso respecto al estimado original de la PO. **Ingreso Real** marcado = "
        "ya recepcionado en CD; si no, es la estimacion vigente mientras sigue en transito."
    )
    if df.empty:
        st.success("No hay compras para mostrar con los filtros actuales.")
        return

    display = _select_cols(df, _COLS_ALMACEN).sort_values("DIAS_DEMORA_ALMACEN", ascending=False, na_position="last")
    st.dataframe(
        display, column_config=_COL_CONFIG, use_container_width=True, hide_index=True,
        height=min(len(display) * 35 + 60, 600),
    )
    n_atraso = int((display["DIAS_DEMORA_ALMACEN"] > 0).sum())
    st.caption(f"{len(display)} lineas | {n_atraso} con atraso de ingreso a almacen")
    download_buttons(display, prefix="cumplimiento_almacen")


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def render_cumplimiento_comex(conn) -> None:
    """Entry point — called from app.py."""
    st.html(page_header(
        "Cumplimiento Comex",
        "Dias de demora de embarque (ETD) e ingreso a almacen vs fechas comprometidas en la PO",
    ))

    if conn is None:
        st.warning("No hay conexion a Snowflake.")
        return

    try:
        df_raw = _load(conn)
    except Exception as e:
        st.error(f"Error al cargar datos: {e}")
        import traceback
        st.code(traceback.format_exc())
        return

    df_raw = apply_pm_filter(df_raw)

    if df_raw.empty:
        st.info("No se encontraron compras de comex en transito o ya ingresadas.")
        return

    df = _enrich(df_raw)

    # ── VP acumulada por SKU (quiebre CD) ────────────────────────────────
    demand_start, demand_end, dias_ventana = _demand_window_for_today()
    try:
        df_vp = _load_vp_acum(id(conn), demand_start, demand_end, dias_ventana, _conn=conn)
        if not df_vp.empty and "SKU_PRODUCTO" in df_vp.columns:
            df = df.merge(
                df_vp[["SKU_PRODUCTO", "VP_ACUMULADA_SOLES"]].drop_duplicates("SKU_PRODUCTO"),
                on="SKU_PRODUCTO",
                how="left",
            )
        else:
            df["VP_ACUMULADA_SOLES"] = np.nan
    except Exception:
        df["VP_ACUMULADA_SOLES"] = np.nan
    df["VP_ACUMULADA_SOLES"] = pd.to_numeric(df["VP_ACUMULADA_SOLES"], errors="coerce").fillna(0)

    # ── Filtros ──────────────────────────────────────────────────────────
    with st.expander("Filtros", expanded=True):
        fc1, fc2, fc3 = st.columns(3)
        with fc1:
            sel_status = st.multiselect(
                "Status", ["En Transito", "Ya Ingreso"], default=[], key="cumplcomex_status",
            )
        with fc2:
            sel_estado = st.multiselect(
                "Estado Importacion",
                sorted(df["ESTADO_IMPORTACION"].dropna().unique().tolist()),
                default=[], key="cumplcomex_estado",
            )
        with fc3:
            sel_proveedor = st.multiselect(
                "Proveedor", sorted(df["PROVEEDOR"].dropna().unique().tolist()),
                default=[], key="cumplcomex_proveedor",
            )

        fc4, fc5, fc6 = st.columns(3)
        with fc4:
            sel_area = st.multiselect(
                "Area", sorted(df["AREA"].dropna().unique().tolist()), default=[], key="cumplcomex_area",
            )
        with fc5:
            sel_linea = st.multiselect(
                "Linea", sorted(df["LINEA"].dropna().unique().tolist()), default=[], key="cumplcomex_linea",
            )
        with fc6:
            sel_marca = st.multiselect(
                "Marca", sorted(df["MARCA"].dropna().unique().tolist()), default=[], key="cumplcomex_marca",
            )
        sel_sublinea = st.multiselect(
            "Sublinea", sorted(df["SUBLINEA"].dropna().unique().tolist()), default=[], key="cumplcomex_sublinea",
        )

    df_f = _apply_filters(df, sel_status, sel_estado, sel_proveedor, sel_area, sel_linea, sel_sublinea, sel_marca)

    # ── Tabs ─────────────────────────────────────────────────────────────
    tab_resumen, tab_embarque, tab_almacen = st.tabs([
        "Resumen",
        f"Cumplimiento Embarque ({len(df_f)})",
        f"Cumplimiento Ingreso Almacen ({len(df_f)})",
    ])

    with tab_resumen:
        _render_resumen(df_f)
    with tab_embarque:
        _render_tabla_embarque(df_f)
    with tab_almacen:
        _render_tabla_almacen(df_f)
